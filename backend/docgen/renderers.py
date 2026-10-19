"""Renderers — one spec in, one real, natively-formatted file out.

Each renderer uses the library idiomatic for its format instead of a
shared "write paragraphs" helper: python-docx with real styles/tables/
headers, ReportLab Platypus for flowing paginated PDF, python-pptx with
per-slide layouts and speaker notes, openpyxl with real worksheet
tables. This is what makes each output *behave* like its format instead
of being the same text dumped into a different extension.
"""
import csv
import re
from pathlib import Path

from docgen.spec import ContentSpec, SlideDeck, Workbook


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------

def render_docx(spec: ContentSpec, path: Path) -> None:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    doc.add_heading(spec.title, level=0)
    if spec.subtitle:
        sub = doc.add_paragraph(spec.subtitle)
        run = sub.runs[0]
        run.italic = True
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    for sec in spec.sections:
        if sec.heading:
            doc.add_heading(sec.heading, level=min(max(sec.level, 1), 2))
        for para in sec.paragraphs:
            p = doc.add_paragraph(para)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if sec.bullets:
            list_style = "List Number" if sec.numbered else "List Bullet"
            for b in sec.bullets:
                doc.add_paragraph(b, style=list_style)
        if sec.table and sec.table.headers:
            _add_docx_table(doc, sec.table)
        if sec.code and sec.code.content:
            _add_docx_code_block(doc, sec.code)

    _add_docx_footer_page_number(doc)
    doc.save(path)


def _add_docx_table(doc, table) -> None:
    n_cols = len(table.headers)
    t = doc.add_table(rows=1, cols=n_cols)
    t.style = "Light Grid Accent 1"
    hdr = t.rows[0].cells
    for i, h in enumerate(table.headers):
        hdr[i].text = h
        for run in hdr[i].paragraphs[0].runs:
            run.bold = True
    for row in table.rows:
        cells = t.add_row().cells
        for i, val in enumerate(row[:n_cols]):
            cells[i].text = str(val)
    doc.add_paragraph()


def _add_docx_code_block(doc, code) -> None:
    from docx.shared import Pt
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    if code.caption:
        cap = doc.add_paragraph()
        run = cap.add_run(code.caption)
        run.bold = True
        run.font.size = Pt(9.5)

    tbl = doc.add_table(rows=1, cols=1)
    cell = tbl.rows[0].cells[0]
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F2F2")
    cell._tc.get_or_add_tcPr().append(shd)

    p = cell.paragraphs[0]
    lines = code.content.splitlines() or [""]
    for i, line in enumerate(lines):
        if i > 0:
            p.add_run().add_break()
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
    doc.add_paragraph()


def _add_docx_footer_page_number(doc) -> None:
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    footer = doc.sections[0].footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Page ")
    run = p.add_run()

    def _field(field_type):
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), field_type)
        return el

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"

    run._r.append(_field("begin"))
    run._r.append(instr)
    run._r.append(_field("separate"))
    run._r.append(_field("end"))


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------

def render_pdf(spec: ContentSpec, path: Path) -> None:
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.lib.enums import TA_LEFT, TA_JUSTIFY
    from reportlab.lib import colors
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem,
        HRFlowable, KeepTogether,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("DocTitle", parent=styles["Title"], fontSize=22, leading=26, spaceAfter=6, alignment=TA_LEFT)
    subtitle_style = ParagraphStyle("DocSubtitle", parent=styles["Normal"], fontSize=12, textColor=colors.HexColor("#595959"), spaceAfter=14)
    h1 = ParagraphStyle("H1", parent=styles["Heading1"], fontSize=15, leading=19, spaceBefore=16, spaceAfter=8, textColor=colors.HexColor("#1F4E78"))
    h2 = ParagraphStyle("H2", parent=styles["Heading2"], fontSize=12.5, leading=16, spaceBefore=10, spaceAfter=6, textColor=colors.HexColor("#2E74B5"))
    body = ParagraphStyle("Body", parent=styles["BodyText"], fontSize=10.5, leading=15.5, spaceAfter=8, alignment=TA_JUSTIFY)
    bullet_style = ParagraphStyle("Bullet", parent=body, spaceAfter=4, alignment=TA_LEFT)
    code_style = ParagraphStyle("Code", parent=styles["Code"], fontSize=8.5, leading=11)

    def esc(t):
        return (t or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    story = [Paragraph(esc(spec.title), title_style)]
    if spec.subtitle:
        story.append(Paragraph(esc(spec.subtitle), subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#BFBFBF"), spaceAfter=10))

    for sec in spec.sections:
        block = []
        if sec.heading:
            block.append(Paragraph(esc(sec.heading), h1 if sec.level == 1 else h2))
        for para in sec.paragraphs:
            block.append(Paragraph(esc(para), body))
        if sec.bullets:
            block.append(ListFlowable(
                [ListItem(Paragraph(esc(b), bullet_style)) for b in sec.bullets],
                bulletType="1" if sec.numbered else "bullet", leftIndent=18, spaceBefore=2, spaceAfter=8,
            ))
        if sec.table and sec.table.headers:
            block.append(_pdf_table(sec.table))
        if sec.code and sec.code.content:
            block.append(_pdf_code_block(sec.code, code_style))
        if not block:
            continue
        # Keep the heading glued to its first bit of content so headings
        # never get orphaned alone at the bottom of a page.
        story.append(KeepTogether(block[:2]) if len(block) >= 2 else block[0])
        story.extend(block[2:])

    def _footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont("Helvetica", 8.5)
        canvas.setFillColor(colors.HexColor("#808080"))
        canvas.drawCentredString(LETTER[0] / 2, 0.45 * inch, f"Page {doc_.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(
        str(path), pagesize=LETTER,
        leftMargin=0.9 * inch, rightMargin=0.9 * inch, topMargin=0.85 * inch, bottomMargin=0.85 * inch,
        title=spec.title,
    )
    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)


def _pdf_table(table):
    from reportlab.platypus import Table, TableStyle, Paragraph
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet

    styles = getSampleStyleSheet()
    cell_style = styles["BodyText"].clone("TblCell")
    cell_style.fontSize = 9
    cell_style.leading = 12
    header_style = styles["BodyText"].clone("TblHeader")
    header_style.textColor = colors.white
    header_style.fontSize = 9

    data = [[Paragraph(f"<b>{h}</b>", header_style) for h in table.headers]]
    for row in table.rows:
        data.append([Paragraph(str(c), cell_style) for c in row])

    t = Table(data, hAlign="LEFT", repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#BFBFBF")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    return t


def _pdf_code_block(code, code_style):
    from reportlab.platypus import Table, TableStyle, Preformatted
    from reportlab.lib import colors

    pre = Preformatted(code.content, code_style)
    t = Table([[pre]], colWidths=[6.3 * 72])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F2F2F2")),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9D9D9")),
        ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    return t


# --------------------------------------------------------------------------
# PPTX
# --------------------------------------------------------------------------

def render_pptx(deck: SlideDeck, path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Pt

    prs = Presentation()
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = deck.title[:100]
    if len(title_slide.placeholders) > 1:
        title_slide.placeholders[1].text = deck.subtitle[:200] or "Generated locally by LOQ AI Workspace"

    bullet_layout = prs.slide_layouts[1]
    for sl in deck.slides:
        slide = prs.slides.add_slide(bullet_layout)
        slide.shapes.title.text = (sl.title or deck.title)[:90]

        body = slide.placeholders[1].text_frame
        body.word_wrap = True
        points = sl.bullets or ["(no content planned)"]
        font_size = 12 if sl.code else (20 if len(points) <= 4 else 16 if len(points) <= 7 else 13)

        for i, text in enumerate(points):
            p = body.paragraphs[0] if i == 0 else body.add_paragraph()
            p.text = str(text)[:220]
            for run in p.runs:
                run.font.size = Pt(font_size)
                if sl.code:
                    run.font.name = "Consolas"

        if sl.notes:
            slide.notes_slide.notes_text_frame.text = sl.notes

    prs.save(path)


# --------------------------------------------------------------------------
# XLSX / CSV
# --------------------------------------------------------------------------

def _coerce_number(val: str):
    s = str(val).strip()
    if re.fullmatch(r"-?\d+", s):
        return int(s)
    if re.fullmatch(r"-?\d+\.\d+", s):
        return float(s)
    return val


def render_xlsx(wb_spec: Workbook, path: Path) -> None:
    from openpyxl import Workbook as XLWorkbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table as XLTable, TableStyleInfo

    wb = XLWorkbook()
    wb.remove(wb.active)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    used_names = set()
    for si, sheet in enumerate(wb_spec.sheets):
        name = re.sub(r"[\\/*?:\[\]]", "", sheet.name)[:31] or f"Sheet{si + 1}"
        base_name, n = name, 2
        while name in used_names:
            name = f"{base_name[:28]}-{n}"
            n += 1
        used_names.add(name)

        ws = wb.create_sheet(title=name)
        n_cols = len(sheet.headers)
        for c, h in enumerate(sheet.headers, start=1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = border
        for r, row in enumerate(sheet.rows, start=2):
            for c, val in enumerate(row[:n_cols], start=1):
                cell = ws.cell(row=r, column=c, value=_coerce_number(val))
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        widths = [len(h) for h in sheet.headers]
        for row in sheet.rows:
            for c, val in enumerate(row[:n_cols]):
                widths[c] = max(widths[c], min(len(str(val)), 60))
        for c, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(c)].width = max(10, min(w + 4, 50))

        ws.freeze_panes = "A2"
        if n_cols and sheet.rows:
            ref = f"A1:{get_column_letter(n_cols)}{len(sheet.rows) + 1}"
            tbl = XLTable(displayName=f"Table{si + 1}", ref=ref)
            tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
            ws.add_table(tbl)

    if not wb.sheetnames:
        wb.create_sheet("Sheet1")
    wb.save(path)


def render_csv(wb_spec: Workbook, path: Path) -> None:
    sheet = next((s for s in wb_spec.sheets if s.headers and s.rows), wb_spec.sheets[0] if wb_spec.sheets else None)
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        if sheet:
            if sheet.headers:
                writer.writerow(sheet.headers)
            for row in sheet.rows:
                writer.writerow(row)
        else:
            writer.writerow([wb_spec.title])


# --------------------------------------------------------------------------
# MD / TXT / HTML
# --------------------------------------------------------------------------

def render_md(spec: ContentSpec) -> str:
    lines = [f"# {spec.title}"]
    if spec.subtitle:
        lines.append(f"*{spec.subtitle}*")
    lines.append("")
    for sec in spec.sections:
        if sec.heading:
            lines.append(f"{'##' if sec.level == 1 else '###'} {sec.heading}")
            lines.append("")
        for para in sec.paragraphs:
            lines.append(para)
            lines.append("")
        if sec.bullets:
            for i, b in enumerate(sec.bullets):
                lines.append(f"{i + 1}. {b}" if sec.numbered else f"- {b}")
            lines.append("")
        if sec.table and sec.table.headers:
            lines.append("| " + " | ".join(sec.table.headers) + " |")
            lines.append("| " + " | ".join(["---"] * len(sec.table.headers)) + " |")
            for row in sec.table.rows:
                lines.append("| " + " | ".join(str(c) for c in row) + " |")
            lines.append("")
        if sec.code and sec.code.content:
            lines.append(f"```{sec.code.language}")
            lines.append(sec.code.content)
            lines.append("```")
            lines.append("")
    return "\n".join(lines).strip() + "\n"


def render_txt(spec: ContentSpec) -> str:
    lines = [spec.title, "=" * len(spec.title)]
    if spec.subtitle:
        lines.append(spec.subtitle)
    lines.append("")
    for sec in spec.sections:
        if sec.heading:
            lines.append(sec.heading)
            lines.append("-" * len(sec.heading))
        for para in sec.paragraphs:
            lines.append(para)
            lines.append("")
        for i, b in enumerate(sec.bullets):
            lines.append(f"{i + 1}. {b}" if sec.numbered else f"  * {b}")
        if sec.bullets:
            lines.append("")
        if sec.table and sec.table.headers:
            widths = [len(h) for h in sec.table.headers]
            for row in sec.table.rows:
                for i, c in enumerate(row):
                    widths[i] = max(widths[i], len(str(c)))

            def fmt_row(vals):
                return "  ".join(str(v).ljust(widths[i]) for i, v in enumerate(vals))

            lines.append(fmt_row(sec.table.headers))
            lines.append("  ".join("-" * w for w in widths))
            for row in sec.table.rows:
                lines.append(fmt_row(row))
            lines.append("")
        if sec.code and sec.code.content:
            lines.append(f"--- {sec.code.caption or 'code'} ---")
            lines.append(sec.code.content)
            lines.append("")
    return "\n".join(lines).strip() + "\n"


def render_html(spec: ContentSpec) -> str:
    def esc(t):
        return (t or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    parts = [
        "<!DOCTYPE html><html><head><meta charset='utf-8'>",
        f"<title>{esc(spec.title)}</title>",
        "<style>"
        "body{font-family:-apple-system,Segoe UI,Roboto,sans-serif;max-width:800px;margin:2.5rem auto;"
        "padding:0 1.25rem;line-height:1.65;color:#1a1a1a}"
        "h1{font-size:2rem;margin-bottom:.2rem}"
        ".subtitle{color:#666;font-size:1.05rem;margin-top:0}"
        "h2{font-size:1.35rem;margin-top:2.2rem;color:#1F4E78;border-bottom:2px solid #eee;padding-bottom:.3rem}"
        "h3{font-size:1.1rem;margin-top:1.6rem;color:#2E74B5}"
        "table{border-collapse:collapse;width:100%;margin:1rem 0}"
        "th,td{border:1px solid #ddd;padding:.5rem .7rem;text-align:left;font-size:.92rem}"
        "th{background:#1F4E78;color:#fff}"
        "tr:nth-child(even){background:#f7f7f7}"
        "pre{background:#f2f2f2;border:1px solid #ddd;border-radius:6px;padding:.9rem 1rem;overflow-x:auto;font-size:.85rem}"
        "code{font-family:Consolas,Menlo,monospace}"
        "li{margin-bottom:.35rem}</style></head><body>",
        f"<h1>{esc(spec.title)}</h1>",
    ]
    if spec.subtitle:
        parts.append(f"<p class='subtitle'>{esc(spec.subtitle)}</p>")
    for sec in spec.sections:
        if sec.heading:
            tag = "h2" if sec.level == 1 else "h3"
            parts.append(f"<{tag}>{esc(sec.heading)}</{tag}>")
        for para in sec.paragraphs:
            parts.append(f"<p>{esc(para)}</p>")
        if sec.bullets:
            tag = "ol" if sec.numbered else "ul"
            parts.append(f"<{tag}>" + "".join(f"<li>{esc(b)}</li>" for b in sec.bullets) + f"</{tag}>")
        if sec.table and sec.table.headers:
            parts.append("<table><thead><tr>" + "".join(f"<th>{esc(h)}</th>" for h in sec.table.headers) + "</tr></thead><tbody>")
            for row in sec.table.rows:
                parts.append("<tr>" + "".join(f"<td>{esc(str(c))}</td>" for c in row) + "</tr>")
            parts.append("</tbody></table>")
        if sec.code and sec.code.content:
            parts.append(f"<pre><code>{esc(sec.code.content)}</code></pre>")
    parts.append("</body></html>")
    return "\n".join(parts)


TEXT_RENDERERS = {"md": render_md, "txt": render_txt, "html": render_html}


def render_document(spec: ContentSpec, fmt: str, out_path: Path) -> None:
    if fmt == "docx":
        render_docx(spec, out_path)
    elif fmt == "pdf":
        render_pdf(spec, out_path)
    else:
        out_path.write_text(TEXT_RENDERERS[fmt](spec), encoding="utf-8")
