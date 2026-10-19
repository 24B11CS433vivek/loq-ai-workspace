"""Turn uploads + pasted text into (prose, code_blocks).

Source code files are never routed through the LLM planner as prose to
be "summarized" — that's how the old pipeline corrupted code (the model
would paraphrase it into prose bullet points). Instead code is preserved
byte-for-byte here and attached to the final document as its own
section/slide/sheet after planning. Only prose (pasted text, and text
extracted from .txt/.md/.docx/.pdf/.pptx/.xlsx/.csv) is handed to the
planner for summarization/structuring.
"""
import io

from fastapi import HTTPException

CODE_LANGUAGES = {
    ".py": "python", ".js": "javascript", ".ts": "typescript",
    ".jsx": "jsx", ".tsx": "tsx", ".java": "java", ".c": "c",
    ".cpp": "cpp", ".go": "go", ".rs": "rust", ".php": "php",
    ".sql": "sql", ".css": "css", ".html": "html",
    ".json": "json", ".xml": "xml", ".yaml": "yaml", ".yml": "yaml",
}
PLAIN_TEXT_EXTS = {".txt", ".md", ".csv"}
OFFICE_EXTS = {".pdf", ".docx", ".xlsx", ".pptx"}

SUPPORTED_UPLOAD_EXTS = set(CODE_LANGUAGES) | PLAIN_TEXT_EXTS | OFFICE_EXTS


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    return "\n\n".join((p.extract_text() or "") for p in reader.pages)


def _extract_docx(data: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for t in doc.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def _extract_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), data_only=True)
    chunks = []
    for name in wb.sheetnames:
        ws = wb[name]
        chunks.append(f"--- Sheet: {name} ---")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                chunks.append(" | ".join(cells))
    return "\n".join(chunks)


def _extract_pptx(data: bytes) -> str:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    chunks = []
    for i, slide in enumerate(prs.slides, start=1):
        chunks.append(f"--- Slide {i} ---")
        for shape in slide.shapes:
            if shape.has_text_frame:
                chunks.append(shape.text_frame.text)
    return "\n".join(chunks)


def _decode_text(data: bytes) -> str:
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


OFFICE_EXTRACTORS = {
    ".pdf": _extract_pdf, ".docx": _extract_docx,
    ".xlsx": _extract_xlsx, ".pptx": _extract_pptx,
}


def load_upload(filename: str, data: bytes) -> dict:
    """Returns {filename, kind: 'code'|'prose', language, text}."""
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in SUPPORTED_UPLOAD_EXTS:
        raise HTTPException(status_code=400, detail=f"Unsupported file type: {ext or 'unknown'}")
    try:
        text = OFFICE_EXTRACTORS[ext](data) if ext in OFFICE_EXTRACTORS else _decode_text(data)
    except Exception as e:
        raise HTTPException(status_code=422, detail=f"Failed to extract {filename}: {e}")

    if ext in CODE_LANGUAGES:
        return {"filename": filename, "kind": "code", "language": CODE_LANGUAGES[ext], "text": text}
    return {"filename": filename, "kind": "prose", "language": "", "text": text}


def gather_sources(pasted_text: str, uploads: list) -> tuple:
    """uploads: list of dicts from load_upload().
    Returns (prose_text_for_planner, code_blocks[{filename,language,content}])."""
    prose_chunks = []
    code_blocks = []
    if pasted_text and pasted_text.strip():
        prose_chunks.append(f"--- Pasted content ---\n{pasted_text.strip()}")
    for u in uploads:
        if u["kind"] == "code":
            code_blocks.append({"filename": u["filename"], "language": u["language"], "content": u["text"]})
        elif u["text"].strip():
            prose_chunks.append(f"--- {u['filename']} ---\n{u['text'].strip()}")
    return "\n\n".join(prose_chunks), code_blocks
