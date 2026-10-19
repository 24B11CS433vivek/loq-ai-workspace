import io
from fastapi import APIRouter, UploadFile, File, HTTPException

router = APIRouter(prefix="/tools/document", tags=["documents"])


def _extract_pdf(data: bytes):
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    pages = [p.extract_text() or "" for p in reader.pages]
    return "\n\n".join(pages), {"pages": len(reader.pages)}


def _extract_docx(data: bytes):
    from docx import Document
    doc = Document(io.BytesIO(data))
    paras = [p.text for p in doc.paragraphs]
    return "\n".join(paras), {"paragraphs": len(paras)}


def _extract_xlsx(data: bytes):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), data_only=True)
    chunks = []
    for name in wb.sheetnames:
        ws = wb[name]
        chunks.append(f"--- Sheet: {name} ---")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                chunks.append(" | ".join(cells))
    return "\n".join(chunks), {"sheets": len(wb.sheetnames)}


def _extract_pptx(data: bytes):
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    chunks = []
    for i, slide in enumerate(prs.slides, start=1):
        chunks.append(f"--- Slide {i} ---")
        for shape in slide.shapes:
            if shape.has_text_frame:
                chunks.append(shape.text_frame.text)
    return "\n".join(chunks), {"slides": len(prs.slides)}


EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".xlsx": _extract_xlsx,
    ".pptx": _extract_pptx,
}


@router.post("/extract")
async def extract_document(file: UploadFile = File(...)):
    filename = file.filename or "upload"
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    extractor = EXTRACTORS.get(ext)
    if not extractor:
        raise HTTPException(status_code=400, detail=f"Unsupported file type: {ext or 'unknown'}")

    data = await file.read()
    try:
        text, meta = extractor(data)
    except Exception as e:
        raise HTTPException(status_code=422, detail=f"Failed to extract {filename}: {e}")

    return {"filename": filename, "text": text, "meta": meta}
